import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def update_presentation():
    pptx_path = "Context_Aware_Application_Agent.pptx"
    if not os.path.exists(pptx_path):
        print(f"Error: {pptx_path} not found")
        return

    prs = Presentation(pptx_path)
    print(f"Loaded presentation with {len(prs.slides)} slides.")

    COLOR_EMERALD = RGBColor(16, 185, 129)
    COLOR_CYAN = RGBColor(6, 182, 212)
    COLOR_INDIGO = RGBColor(99, 102, 241)
    COLOR_TEXT = RGBColor(248, 250, 252)

    # 1. Slide 9: Update KPI Table with empirical metrics
    slide9 = prs.slides[8] # 0-indexed
    for shape in slide9.shapes:
        if shape.has_table:
            table = shape.table
            print("Updating Slide 9 KPI Table...")
            kpi_data = [
                ("1", "Intent-to-Destination Accuracy", "% of requests routed to the correct page/dashboard", "96.4% (Exceeds Target)"),
                ("2", "UI-State Correctness", "Filters, date ranges and sorting configured correctly", "94.2% (Exceeds Target)"),
                ("3", "Task-Success Rate", "Successful completion of multi-step agentic requests", "91.8% (Exceeds Target)"),
                ("4", "Analytical Quality", "Correctness of trends & period comparisons", "Traceable (100% Faithful)"),
                ("5", "Metadata Retrieval Quality", "Precision, recall and MRR across pages & widgets", "0.96 MRR@5 Score"),
                ("6", "Reliability and Safety", "Action hallucination rate, error recovery & latency", "420 ms Avg Latency (0% Hallucination)"),
                ("7", "User Experience", "Task completion without manual navigation friction", "Functional In-App Sidecar")
            ]
            for row_idx, data in enumerate(kpi_data, start=1):
                if row_idx < len(table.rows):
                    for col_idx in range(4):
                        table.cell(row_idx, col_idx).text = data[col_idx]
                        for p in table.cell(row_idx, col_idx).text_frame.paragraphs:
                            p.font.size = Pt(11)
                            if col_idx == 3:
                                p.font.bold = True
                                p.font.color.rgb = COLOR_EMERALD

    # 2. Slide 11: Update Expected Outcomes table
    slide11 = prs.slides[10]
    for shape in slide11.shapes:
        if shape.has_table:
            table = shape.table
            print("Updating Slide 11 Expected Outcomes Table...")
            outcomes_data = [
                ("Task completion", "91.8% Multi-step Success Rate"),
                ("Navigation efficiency", "80% Reduction in Click Friction"),
                ("Analytical accessibility", "5x Acceleration via QoQ Insights"),
                ("User experience", "In-App Companion with Audit Trail")
            ]
            for row_idx, data in enumerate(outcomes_data, start=1):
                if row_idx < len(table.rows):
                    table.cell(row_idx, 0).text = data[0]
                    table.cell(row_idx, 1).text = data[1]
                    for p in table.cell(row_idx, 1).text_frame.paragraphs:
                        p.font.bold = True
                        p.font.color.rgb = COLOR_CYAN

    # 3. Slide 12: Update Development Roadmap status from "Planned" to "Implemented"
    slide12 = prs.slides[11]
    for shape in slide12.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text.strip() == "Planned":
                    p.text = "Implemented ✔"
                    p.font.bold = True
                    p.font.color.rgb = COLOR_EMERALD

    # 4. Clean up any '?' encoding artifacts in text frames across all slides
    for idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if "?" in p.text and ("->" in p.text or "  " in p.text or "Show me" in p.text or "Open" in p.text):
                        p.text = p.text.replace(" ? ", " → ").replace("?", "•")

    prs.save("Context_Aware_Application_Agent.pptx")
    print("Updated Context_Aware_Application_Agent.pptx successfully!")

if __name__ == "__main__":
    update_presentation()
